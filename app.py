"""
Веб-приложение для пользователей: только читает готовую базу base/hints.db.
Запуск:  uvicorn app:app --host 0.0.0.0 --port 8000
"""
import io, json, os, sqlite3
from datetime import date
from pathlib import Path
from urllib.parse import quote
from fastapi import Body, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse

from updater import Updater, UpdateError

BASE = Path(__file__).parent
DB = BASE / "base" / "hints.db"
app = FastAPI(title="Помощник валидации проверок")
updater = Updater(BASE)


SCHEMA_VERSION = 1   # какой формат базы понимает приложение (см. builder/precompute.py)


def base_meta() -> dict:
    """Сведения о базе: дата сборки, число проверок; ok=False — базы нет или формат чужой."""
    if not DB.exists():
        return {"ok": False, "reason": "нет базы base/hints.db"}
    try:
        con = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
        try:
            meta = dict(con.execute("SELECT key, value FROM meta"))
        finally:
            con.close()
    except sqlite3.Error:
        meta = {}
    v = int(meta.get("schema_version") or 0)
    if v != SCHEMA_VERSION:
        return {"ok": False, **meta, "reason": f"формат базы {v or 'старый'}, приложение понимает {SCHEMA_VERSION} — обновите программу"}
    return {"ok": True, **meta}


def q(sql, *args):
    m = base_meta()
    if not m["ok"]:
        raise HTTPException(503, m["reason"])
    con = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in con.execute(sql, args)]
    finally:
        con.close()


@app.get("/")
def index():
    return FileResponse(BASE / "static" / "index.html")


@app.get("/api/meta")
def meta():
    return base_meta()


@app.get("/api/checks")
def checks(search: str = ""):
    rows = q("SELECT id, number, title, tags, opened FROM checks ORDER BY CAST(number AS INTEGER) DESC")
    s = search.lower()
    out = []
    for r in rows:
        if s and s not in (r["number"] + " " + r["title"]).lower():
            continue
        r["tags"] = json.loads(r["tags"])
        out.append(r)
    return out


@app.get("/api/checks/{check_id}")
def check(check_id: str):
    rows = q("SELECT * FROM checks WHERE id=?", check_id)
    if not rows:
        raise HTTPException(404, "Проверка не найдена")
    c = rows[0]
    c["tags"] = json.loads(c["tags"])
    c["items"] = json.loads(c["items"])
    info = {r["id"]: r for r in q("SELECT id, number, title FROM checks")}
    c["similar"] = [
        {**r, "number": info[r["similar_id"]]["number"], "title": info[r["similar_id"]]["title"],
         "common_tags": json.loads(r["common_tags"])}
        for r in q("SELECT similar_id, score, common_tags, reason FROM similar WHERE check_id=? ORDER BY score DESC", check_id)
    ]
    reasons = {s["similar_id"]: s["reason"] for s in c["similar"]}
    hints = q("SELECT * FROM hints WHERE check_id=? ORDER BY hint_no", check_id)
    for h in hints:
        h["sources"] = json.loads(h["sources"])
        for s in h["sources"]:
            s["title"] = info[s["check_id"]]["title"]
            s["reason"] = reasons.get(s["check_id"], "")
    c["hints"] = hints
    return c


# ---------------- Экспорт ----------------

def _export_data(check_id, done):
    c = check(check_id)
    done_set = {int(x) for x in done.split(",") if x.strip().isdigit()}
    return c, done_set


def _src_text(h):
    return "; ".join(f"№{s['number']} «{s['title']}» ({round(s['score']*100)}%)" for s in h["sources"])


def _details(h, sep="\n"):
    """Результаты / документы / методы из источников."""
    res, inv, meth = [], [], []
    for s in h["sources"]:
        for it in s["items"]:
            if it.get("results"):
                res.append(f"№{s['number']}: {it['results']}")
            if it.get("invoice"):
                inv.append(it["invoice"])
            if it.get("method"):
                meth.append(it["method"].replace("\n", "; "))
    u = lambda xs: list(dict.fromkeys(xs))
    return sep.join(u(res)), sep.join(u(inv)), sep.join(u(meth))


def _status(h):
    t = f"нарушение подтверждено в {h['confirmed']} из {h['total']}" if h["confirmed"] else "не подтверждалось"
    return t


def _file(buf, name, mime):
    buf.seek(0)
    return StreamingResponse(buf, media_type=mime, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@app.get("/api/checks/{check_id}/export.xlsx")
def export_xlsx(check_id: str, done: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    c, done_set = _export_data(check_id, done)

    wb = Workbook()
    ws = wb.active
    ws.title = "Чек-лист"
    ws["A1"] = f"Чек-лист валидации: {c['title']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Проверка №{c['number']} · сформирован {date.today():%d.%m.%Y}"
    ws["A3"] = "Суть: " + ", ".join(c["tags"])
    head = ["№", "Задача для проверки", "Гипотеза", "Есть в программе", "Итог в похожих",
            "Отметка", "Комментарий валидатора", "Что выявляли", "Документы", "Методы", "Источники", "Типовая"]
    ws.append([])
    ws.append(head)
    hdr_row = ws.max_row
    fill = PatternFill("solid", fgColor="2456D6")
    thin = Side(style="thin", color="D0D5DD")
    for cell in ws[hdr_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for h in c["hints"]:
        res, inv, meth = _details(h)
        ws.append([h["hint_no"], h["task"], h["hypo"], h["covered"] or "нет", _status(h),
                   "✓" if h["hint_no"] in done_set else "", "", res, inv, meth, _src_text(h),
                   "да" if h["typical"] else ""])
    for row in ws.iter_rows(min_row=hdr_row, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if cell.row > hdr_row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, w in zip("ABCDEFGHIJKL", [5, 50, 35, 30, 22, 9, 35, 50, 30, 25, 40, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

    ws2 = wb.create_sheet("Похожие проверки")
    ws2.append(["Проверка", "Название", "Похожесть", "Почему похожа"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = fill
    for s in c["similar"]:
        ws2.append([f"№{s['number']}", s["title"], f"{round(s['score']*100)}%", s["reason"]])
    for col, w in zip("ABCD", [12, 40, 11, 80]):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO(); wb.save(buf)
    return _file(buf, f"Чек-лист №{c['number']}.xlsx",
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/checks/{check_id}/export.docx")
def export_docx(check_id: str, done: str = ""):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    c, done_set = _export_data(check_id, done)

    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
    d.add_heading(f"Чек-лист валидации: {c['title']}", level=1)
    p = d.add_paragraph()
    r = p.add_run(f"Проверка №{c['number']} · сформирован {date.today():%d.%m.%Y}")
    r.font.color.rgb = RGBColor(0x6B, 0x76, 0x86)
    for label, key in [("Задача программы", "purpose"), ("Метрики", "metrics"), ("Основания", "background")]:
        if c[key]:
            p = d.add_paragraph(); p.add_run(label + ": ").bold = True; p.add_run(c[key])
    d.add_paragraph("Суть проверки: " + ", ".join(c["tags"]))

    d.add_heading("Задачи из похожих проверок", level=2)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, txt in zip(t.rows[0].cells, ["№", "Задача и источники", "✓", "Комментарий"]):
        cell.text = ""
        cell.paragraphs[0].add_run(txt).bold = True
    gray = RGBColor(0x6B, 0x76, 0x86)
    for h in c["hints"]:
        row = t.add_row().cells
        row[0].text = str(h["hint_no"])
        row[1].text = ""
        row[1].paragraphs[0].add_run(h["task"] + (" (типовая)" if h["typical"] else "")).bold = True
        res, inv, meth = _details(h, "; ")
        lines = [("Гипотеза", h["hypo"]), ("Итог в похожих", _status(h)),
                 ("В программе", h["covered"] or "нет"), ("Что выявляли", res),
                 ("Документы", inv), ("Методы", meth), ("Источники", _src_text(h))]
        for k, v in lines:
            if v:
                rr = row[1].add_paragraph().add_run(f"{k}: {v}")
                rr.font.size = Pt(8.5); rr.font.color.rgb = gray
        row[2].text = "✓" if h["hint_no"] in done_set else "☐"
        row[3].text = ""
    t.autofit = False
    widths = [Cm(1), Cm(10.5), Cm(1), Cm(4.5)]
    for col, w in zip(t.columns, widths):
        col.width = w
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    d.add_heading("Похожие проверки", level=2)
    for s in c["similar"]:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(f"№{s['number']} «{s['title']}» — {round(s['score']*100)}%. ").bold = True
        p.add_run(s["reason"])

    buf = io.BytesIO(); d.save(buf)
    return _file(buf, f"Чек-лист №{c['number']}.docx",
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ---------------- Обновление из Bitbucket ----------------
# Проверить версию может кто угодно; ставить обновление и перезапускать —
# только с этого же компьютера (если сервер открыт в сеть через --host 0.0.0.0).

def _local_only(request: Request):
    host = request.client.host if request.client else ""
    if host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        raise UpdateError("E_FORBIDDEN", "Обновлять можно только на компьютере, где запущен сервер",
                          status=403)


@app.exception_handler(UpdateError)
def _update_error(_request, e: UpdateError):
    return JSONResponse(e.as_dict(), status_code=e.status)


@app.get("/api/update/info")
def update_info():
    return updater.info()


@app.get("/api/update/check")
def update_check():
    return updater.check()


@app.post("/api/update/login")
def update_login(request: Request, body: dict = Body(...)):
    _local_only(request)
    updater.set_credentials(body.get("login", ""), body.get("password", ""))
    return updater.check()


@app.post("/api/update/apply")
def update_apply(request: Request):
    _local_only(request)
    return updater.apply()


@app.get("/api/update/status")
def update_status():
    return updater.job


@app.post("/api/update/restart")
def update_restart(request: Request):
    _local_only(request)
    if updater.job.get("state") != "done":
        raise UpdateError("E_UPDATE", "Сначала установите обновление", status=409)
    return updater.restart(os.environ.get("VH_HOST", "127.0.0.1"), int(os.environ.get("VH_PORT", "8000")))
